"""
FML Freight Solutions — Inspection Report Generator
====================================================
Usage:
  python generate_report.py                          # process ALL inspections in config.inspections_dir
  python generate_report.py --ba BA2797              # process a single BA number
  python generate_report.py --config my.yaml         # use a custom config file

Output: one .pptx per inspection written to config.output_dir

Now with MPO image support! Converts MPO (3D camera) files to JPEG automatically.
"""

import argparse
import json
import os
import sys
import glob
import tempfile
from pathlib import Path
from datetime import date

import yaml
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image  # Added for MPO conversion

# ── Brand colours ──────────────────────────────────────────────────────────────
DARK_TEAL = RGBColor(0x00, 0x3D, 0x3B)  # slide background
MID_TEAL = RGBColor(0x00, 0x7A, 0x7A)   # accents
LIGHT_TEAL = RGBColor(0x5F, 0xC4, 0xBF) # secondary text / logo
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF0, 0xF4, 0xF4)
DARK_GREY = RGBColor(0x1A, 0x1A, 0x1A)
MID_GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GREY = RGBColor(0xCC, 0xCC, 0xCC)
RED_DAMAGE = RGBColor(0xCC, 0x22, 0x22)
GREEN_OK = RGBColor(0x22, 0x88, 0x44)
AMBER = RGBColor(0xE6, 0x8A, 0x00)

# ── Slide dimensions (16:9 widescreen) ─────────────────────────────────────────
SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

# ═══════════════════════════════════════════════════════════════════════════════
# Helper utilities
# ═══════════════════════════════════════════════════════════════════════════════

def _set_bg(slide, colour: RGBColor):
    from pptx.oxml.ns import qn
    from lxml import etree
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = colour

def _add_textbox(slide, left, top, width, height, text,
                 font_size=12, bold=False, italic=False,
                 colour=WHITE, align=PP_ALIGN.LEFT,
                 wrap=True, font_name="Calibri"):
    txb = slide.shapes.add_textbox(left, top, width, height)
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic